from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZIP_STORED, ZipFile, ZipInfo

from docx import Document
from docx.document import Document as DocumentType
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.section import Section
from docx.shared import Cm, Pt
from docx.table import Table, _Cell

from riskprobe.analysis_contracts import MetricDistribution
from riskprobe.report_models import ReportCondition, ReportModel, ReportRule

_FIXED_CORE_TIME = datetime(1980, 1, 1, tzinfo=timezone.utc)
_FIXED_ZIP_TIME = (1980, 1, 1, 0, 0, 0)
_FONT_NAME = "Microsoft YaHei"


def render_report_docx(model: ReportModel) -> bytes:
    document = Document()
    _configure_document(document)
    _add_cover(document, model)
    _add_executive_summary(document, model)
    _add_stage_summary(document, model)
    _add_profile_partition(document, model)
    _add_discovery(document, model)
    _add_top_rules(document, model)
    _add_scorecard(document, model)
    _add_institution_stability(document, model)
    _add_decision_review(document, model)
    _add_limitations_appendix(document, model)
    buffer = BytesIO()
    document.save(buffer)
    return _canonicalize_docx_package(buffer.getvalue())


def _configure_document(document: DocumentType) -> None:
    properties = document.core_properties
    properties.title = "RiskProbe 风险分析报告"
    properties.subject = "RiskProbe deterministic terminal report"
    properties.author = ""
    properties.last_modified_by = ""
    properties.keywords = ""
    properties.comments = ""
    properties.created = _FIXED_CORE_TIME
    properties.modified = _FIXED_CORE_TIME

    section = document.sections[0]
    _configure_section(section, landscape=False)
    for style_name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Heading 3"):
        style = document.styles[style_name]
        style.font.name = _FONT_NAME
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), _FONT_NAME)
    document.styles["Normal"].font.size = Pt(10.5)
    document.styles["Title"].font.size = Pt(24)
    document.styles["Heading 1"].font.size = Pt(16)
    document.styles["Heading 2"].font.size = Pt(13)
    document.styles["Heading 3"].font.size = Pt(11)
    _add_footer(section)


def _configure_section(section: Section, *, landscape: bool) -> None:
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Cm(29.7)
        section.page_height = Cm(21)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)


def _add_footer(section: Section) -> None:
    section.footer.is_linked_to_previous = False
    paragraph = section.footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run("RiskProbe · ")
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, end))


def _add_cover(document: DocumentType, model: ReportModel) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run("\n\n\n")
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run(model.title)
    subtitle = document.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("完整分析与 Host 决策终态报告")
    document.add_paragraph("")
    profile = model.profile
    summary = model.analysis_summary
    dataset_id = (
        profile.dataset_id
        if profile is not None
        else summary.input.dataset_id
        if summary is not None and summary.input is not None
        else "N/A"
    )
    grade = (
        profile.metadata_grade
        if profile is not None
        else summary.profile.metadata_grade
        if summary is not None and summary.profile is not None
        else "N/A"
    )
    _add_key_values(
        document,
        (
            ("Dataset", dataset_id),
            ("Run ID", model.run_id or "N/A"),
            ("Session ID", model.session_id or "N/A"),
            ("Scope", model.scope),
            ("Terminal status", model.terminal_status or "N/A"),
            ("Metadata Grade", grade),
        ),
    )
    document.add_paragraph(
        "时间切片稳定性：已评估"
        if model.time_validation_applied
        else "时间切片稳定性：未评估"
    )
    if grade == "B":
        document.add_paragraph(
            "限制：表现窗口未知；Stable 不代表严格 OOT、生产就绪或自动上线。"
        )
    if model.time_validation_applied:
        document.add_paragraph(
            "时间切片评估不等于已知表现窗口、严格 OOT 或生产就绪。"
            if grade == "B"
            else "时间切片评估不等于已知表现窗口或生产就绪。"
        )
    document.add_page_break()


def _add_executive_summary(document: DocumentType, model: ReportModel) -> None:
    document.add_heading("1. 管理摘要", level=1)
    summary = model.analysis_summary
    profile = model.profile
    analysis_profile = summary.profile if summary is not None else None
    discovery = summary.discovery if summary is not None else None
    scorecard = summary.scorecard if summary is not None else None
    _add_key_values(
        document,
        (
            ("样本量", _integer(profile.row_count if profile is not None else analysis_profile.row_count if analysis_profile is not None else None)),
            ("特征数", _integer(profile.feature_count if profile is not None else analysis_profile.feature_count if analysis_profile is not None else None)),
            ("正样本率", _number(profile.positive_rate if profile is not None else analysis_profile.positive_rate if analysis_profile is not None else None)),
            ("候选规则数", _integer(discovery.candidate_rule_count if discovery is not None else None)),
            ("入选规则数", str(len(model.all_rules))),
            ("Stable 规则数", str(model.grade_counts.get("Stable", 0))),
            ("评分卡状态", scorecard.status.value if scorecard is not None else "N/A"),
            ("最终决策", model.decision_summary.decision_status if model.decision_summary is not None else model.terminal_status or "N/A"),
        ),
    )


def _add_stage_summary(document: DocumentType, model: ReportModel) -> None:
    document.add_heading("2. 全流程阶段摘要", level=1)
    summary = model.analysis_summary
    if summary is None:
        document.add_paragraph("未提供结构化阶段摘要。")
        return
    _add_table(
        document,
        ("Stage", "Status", "Enabled", "Output", "Reason", "Duration ms", "Limitations"),
        tuple(
            (
                stage.name.value,
                stage.status.value,
                "是" if stage.enabled else "否",
                "是" if stage.output_available else "否",
                stage.reason_code or "N/A",
                _integer(stage.duration_ms),
                ", ".join(stage.limitations) or "N/A",
            )
            for stage in summary.stages
        ),
    )


def _add_profile_partition(document: DocumentType, model: ReportModel) -> None:
    document.add_heading("3. 数据画像、质量与切分", level=1)
    profile = model.profile
    summary = model.analysis_summary
    analysis_profile = summary.profile if summary is not None else None
    if profile is not None:
        _add_key_values(
            document,
            (
                ("Rows", str(profile.row_count)),
                ("Features", str(profile.feature_count)),
                ("Positive rate", _number(profile.positive_rate)),
                ("Segment count", str(profile.segment_count)),
                ("Snapshot range", f"{profile.snapshot_min.isoformat() if profile.snapshot_min else 'N/A'} to {profile.snapshot_max.isoformat() if profile.snapshot_max else 'N/A'}"),
            ),
        )
        document.add_heading("Quality Issues", level=2)
        _add_table(
            document,
            ("Severity", "Code", "Affected Rows"),
            tuple((item.severity, item.code, item.affected_rows) for item in profile.issues),
        )
    elif analysis_profile is not None:
        _add_key_values(
            document,
            (
                ("Rows", str(analysis_profile.row_count)),
                ("Features", str(analysis_profile.feature_count)),
                ("Numeric features", str(analysis_profile.numeric_feature_count)),
                ("Positive rate", _number(analysis_profile.positive_rate)),
                ("Issue codes", ", ".join(analysis_profile.issue_codes) or "None"),
            ),
        )
    else:
        document.add_paragraph("未提供数据画像。")
    document.add_heading("数据切分", level=2)
    partition = summary.partition if summary is not None else None
    if partition is None:
        document.add_paragraph("未提供切分摘要。")
        return
    _add_table(
        document,
        ("Strategy", "Mode", "Time Requested", "Time Applied", "Train", "Test", "Holdout", "Excluded Null", "Fallback"),
        ((partition.strategy, partition.mode, partition.requested_time_validation, partition.applied_time_validation, partition.train_rows, partition.test_rows, partition.holdout_rows, partition.excluded_null_snapshot_rows, partition.fallback_reason_code or "N/A"),),
    )


def _add_discovery(document: DocumentType, model: ReportModel) -> None:
    document.add_heading("4. 规则发现与筛选", level=1)
    _add_table(
        document,
        ("Stable", "Local", "Unstable", "Suspicious"),
        ((model.grade_counts.get("Stable", 0), model.grade_counts.get("Local", 0), model.grade_counts.get("Unstable", 0), model.grade_counts.get("Suspicious", 0)),),
    )
    summary = model.analysis_summary
    discovery = summary.discovery if summary is not None else None
    if discovery is None:
        document.add_paragraph(f"入选规则：{len(model.all_rules)}")
        return
    _add_key_values(
        document,
        (
            ("Input features", str(discovery.input_feature_count)),
            ("Eligible features", str(discovery.eligible_feature_count)),
            ("Candidate rules", str(discovery.candidate_rule_count)),
            ("Selected rules", str(discovery.selected_rule_count)),
            ("Single candidates/selected", f"{discovery.single_candidate_count}/{discovery.single_rule_count}"),
            ("Pair candidates/selected", f"{discovery.pair_candidate_count}/{discovery.pair_rule_count}"),
            ("Lift", _distribution(discovery.lift)),
            ("Coverage", _distribution(discovery.support)),
            ("Precision", _distribution(discovery.precision)),
        ),
    )


def _add_top_rules(document: DocumentType, model: ReportModel) -> None:
    section = document.add_section(WD_SECTION.NEW_PAGE)
    _configure_section(section, landscape=True)
    _add_footer(section)
    document.add_heading("5. TOP10 规则", level=1)
    headers = (
        "Rank",
        "Rule ID",
        "规则条件",
        "Origin",
        "Grade",
        "Support",
        "Test Lift",
        "Holdout Lift",
        "Coverage",
        "Hit Bad Rate",
        "signed KS",
        "Adjusted p-value",
        "Lift CI",
        "Segment Consistency",
    )
    rows = tuple(
        (
            rule.rank,
            rule.rule_id,
            _rule_text(rule),
            rule.origin,
            rule.grade,
            rule.test.support_count,
            _number(rule.test.lift),
            _number(rule.holdout.lift if rule.holdout is not None else None),
            _number(rule.test.coverage),
            _number(rule.test.hit_bad_rate),
            _number(rule.test.signed_ks),
            _number(rule.adjusted_p_value),
            f"{_number(rule.lift_ci[0])}–{_number(rule.lift_ci[1])}",
            _number(rule.segment_consistency),
        )
        + ((_number(rule.time_decay),) if model.time_validation_applied else ())
        for rule in model.top_rules
    )
    table = _add_table(
        document,
        headers + (("Time Decay",) if model.time_validation_applied else ()),
        rows,
        font_size=6.5,
    )
    if table is not None:
        table.autofit = True
    for rule in model.top_rules:
        document.add_heading(f"#{rule.rank} {rule.rule_id}", level=2)
        items = [
            ("条件", _rule_text(rule)),
            ("Grade / Origin", f"{rule.grade} / {rule.origin}"),
            (
                "Train / Test / Holdout Lift",
                f"{_number(rule.train.lift)} / {_number(rule.test.lift)} / {_number(rule.holdout.lift if rule.holdout is not None else None)}",
            ),
            ("Segment Consistency", _number(rule.segment_consistency)),
        ]
        if model.time_validation_applied:
            items.append(("Time Decay", _number(rule.time_decay)))
        items.append(("Limitations", ", ".join(rule.limitations) or "None"))
        _add_key_values(document, tuple(items))
    portrait = document.add_section(WD_SECTION.NEW_PAGE)
    _configure_section(portrait, landscape=False)
    _add_footer(portrait)


def _add_scorecard(document: DocumentType, model: ReportModel) -> None:
    document.add_heading("6. 评分卡", level=1)
    summary = model.analysis_summary
    scorecard = summary.scorecard if summary is not None else None
    if scorecard is None:
        document.add_paragraph("未提供评分卡摘要。")
        return
    _add_key_values(
        document,
        (
            ("Enabled", str(scorecard.enabled)),
            ("Status", scorecard.status.value),
            ("Reason", scorecard.reason_code or "N/A"),
            ("Model", scorecard.model_type or "N/A"),
            ("Calibrated", str(scorecard.calibrated) if scorecard.calibrated is not None else "N/A"),
            ("Imbalance strategy", scorecard.imbalance_strategy or "N/A"),
            ("Intercept", _number(scorecard.intercept)),
        ),
    )
    _add_table(
        document,
        ("Split", "Status", "Samples", "Positive Rate", "AUC", "KS", "Gini", "Mean Score", "Reason"),
        tuple((split.split, split.status.value, split.sample_count, _number(split.positive_rate), _number(split.auc), _number(split.ks), _number(split.gini), _number(split.mean_score), split.reason_code or "N/A") for split in scorecard.splits),
    )
    document.add_heading("WOE / IV 特征", level=2)
    _add_table(
        document,
        ("Feature", "IV", "Bins", "Monotonic", "Missing Bin"),
        tuple((item.feature.value, _number(item.iv), item.bin_count, item.monotonic, item.has_missing_bin) for item in scorecard.feature_summaries),
    )
    document.add_heading("模型系数", level=2)
    _add_table(
        document,
        ("Term", "Kind", "Coefficient"),
        tuple((term.term.value, term.kind, _number(term.coefficient)) for term in scorecard.terms),
    )


def _add_institution_stability(document: DocumentType, model: ReportModel) -> None:
    document.add_heading("7. 机构与稳定性分析", level=1)
    _add_table(
        document,
        ("Rule ID", "Institution Token", "Support", "Coverage", "Hit Bad Rate", "Lift", "Direction"),
        tuple((row.rule_id, row.institution_token, row.support_count, _number(row.coverage), _number(row.hit_bad_rate), _number(row.lift), row.direction) for row in model.institution_evidence),
    )
    summary = model.institution_summary
    if summary is None:
        document.add_paragraph("未提供机构分析摘要。")
        return
    _add_key_values(
        document,
        (
            ("Eligible institutions", str(summary.eligible_count)),
            ("Triggered local discovery", str(summary.triggered_count)),
            ("Blocked local discovery", str(summary.blocked_count)),
            ("Interpretation", summary.interpretation),
        ),
    )


def _add_decision_review(document: DocumentType, model: ReportModel) -> None:
    document.add_heading("8. 诊断、建议与审核", level=1)
    summary = model.analysis_summary
    diagnostics = summary.diagnostics if summary is not None else None
    if diagnostics is not None:
        _add_key_values(
            document,
            (
                ("Diagnostic time enabled", str(diagnostics.diagnostic_time_enabled)),
                ("Finding counts by kind", ", ".join(f"{key}={value}" for key, value in diagnostics.finding_counts_by_kind.items()) or "None"),
                ("Finding counts by severity", ", ".join(f"{key}={value}" for key, value in diagnostics.finding_counts_by_severity.items()) or "None"),
            ),
        )
    _add_key_values(
        document,
        (
            ("Proposal actions", ", ".join(model.proposal_action_codes) or "None"),
            ("Host decision reasons", ", ".join(model.decision_reason_codes) or "None"),
            ("Diagnosis evidence IDs", ", ".join(model.diagnosis_evidence_ids) or "None"),
            ("Agent tool sequence", " → ".join(model.agent_tool_sequence) or "None"),
            ("Agent state history", " → ".join(model.agent_state_history) or "None"),
            ("Agent summary", model.agent_summary or "N/A"),
        ),
    )
    if model.findings:
        _add_table(
            document,
            ("Evidence ID", "Kind", "Severity", "Summary", "Limitations"),
            tuple((finding.evidence_id, finding.kind, finding.severity, finding.summary, ", ".join(finding.limitations) or "None") for finding in model.findings),
        )
    decision = model.decision_summary
    if decision is None:
        document.add_paragraph(
            "无需 Host proposal/无可执行动作。"
            if model.terminal_status == "no_action"
            else "Host 决策尚未执行。"
        )
    else:
        _add_key_values(
            document,
            (
                ("Decision status", decision.decision_status),
                ("Final status", decision.final_status),
                ("Selected actions", ", ".join(decision.selected_action_codes) or "None"),
                ("Recommendation status", decision.recommendation_status.value),
                ("Review approved", str(decision.review_approved)),
                ("Review reasons", ", ".join(decision.review_reason_codes) or "None"),
                ("Evidence complete", str(decision.evidence_complete)),
                ("Retry count", str(decision.retry_count)),
                ("No action required", str(decision.no_action_required)),
                ("Tool sequence", " → ".join(decision.tool_sequence)),
            ),
        )
        if decision.recommendations:
            document.add_heading("Recommendations", level=2)
            _add_table(
                document,
                ("Action", "Evidence ID", "Parent Findings", "Human Approval", "Analysis Only", "Limitations"),
                tuple((item.action_code, item.evidence_id, ", ".join(item.parent_finding_ids), item.human_approval_required, item.analysis_only, ", ".join(item.limitations) or "None") for item in decision.recommendations),
            )
    if model.review is not None:
        document.add_heading("Review", level=2)
        _add_key_values(
            document,
            (
                ("Approved", str(model.review.approved)),
                ("Reasons", ", ".join(model.review.reason_codes) or "None"),
                ("Evidence IDs", ", ".join(model.review.evidence_ids) or "None"),
                ("Retry allowed", str(model.review.retry_allowed)),
                ("No action required", str(model.review.no_action_required)),
            ),
        )


def _add_limitations_appendix(document: DocumentType, model: ReportModel) -> None:
    document.add_heading("9. 限制", level=1)
    limitations = set(model.limitations)
    limitations.update(item for rule in model.top_rules for item in rule.limitations)
    limitations.add(
        "时间切片稳定性：已评估"
        if model.time_validation_applied
        else "时间切片稳定性：未评估"
    )
    metadata_grade = (
        model.profile.metadata_grade
        if model.profile is not None
        else model.analysis_summary.profile.metadata_grade
        if model.analysis_summary is not None
        and model.analysis_summary.profile is not None
        else None
    )
    if metadata_grade == "B":
        limitations.add("label performance window unknown")
    if model.time_validation_applied:
        limitations.add(
            "时间切片评估不等于已知表现窗口、严格 OOT 或生产就绪"
            if metadata_grade == "B"
            else "时间切片评估不等于已知表现窗口或生产就绪"
        )
    for item in sorted(limitations):
        document.add_paragraph(item, style="List Bullet")
    document.add_heading("指标定义", level=2)
    definitions = [
        "Lift：命中样本坏样本率相对总体坏样本率的倍数。",
        "signed KS：命中坏样本率减命中好样本率，保留风险方向。",
        "Adjusted p-value：多重检验校正后的显著性指标。",
        "Segment Consistency：规则跨机构/分群方向一致性。",
    ]
    if model.time_validation_applied:
        definitions.append("Time Decay：规则跨时间窗口的最大 Lift 衰减。")
    for item in definitions:
        document.add_paragraph(item, style="List Bullet")
    document.add_heading("10. 附录：全部入选规则简表", level=1)
    _add_table(
        document,
        ("Rule ID", "规则条件", "Origin", "Grade", "Test Lift", "Holdout Lift", "Coverage", "signed KS"),
        tuple((rule.rule_id, _rule_text(rule), rule.origin, rule.grade, _number(rule.test.lift), _number(rule.holdout.lift if rule.holdout is not None else None), _number(rule.test.coverage), _number(rule.test.signed_ks)) for rule in model.all_rules),
        font_size=8,
    )


def _add_key_values(document: DocumentType, items: tuple[tuple[str, str], ...]) -> None:
    for key, value in items:
        paragraph = document.add_paragraph()
        key_run = paragraph.add_run(f"{key}: ")
        key_run.bold = True
        paragraph.add_run(value)


def _add_table(
    document: DocumentType,
    headers: tuple[str, ...],
    rows: tuple[tuple[object, ...], ...],
    *,
    font_size: float = 8.5,
) -> Table | None:
    if not headers:
        return None
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.autofit = True
    header_cells = table.rows[0].cells
    for index, header in enumerate(headers):
        _set_cell_text(header_cells[index], header, bold=True, font_size=font_size)
    _repeat_table_header(table)
    if rows:
        for row in rows:
            cells = table.add_row().cells
            for index, value in enumerate(row):
                _set_cell_text(cells[index], str(value), bold=False, font_size=font_size)
    else:
        cells = table.add_row().cells
        _set_cell_text(cells[0], "None", bold=False, font_size=font_size)
        for cell in cells[1:]:
            _set_cell_text(cell, "—", bold=False, font_size=font_size)
    return table


def _set_cell_text(cell: _Cell, value: str, *, bold: bool, font_size: float) -> None:
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(value)
    run.bold = bold
    run.font.name = _FONT_NAME
    run.font.size = Pt(font_size)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), _FONT_NAME)


def _repeat_table_header(table: Table) -> None:
    properties = table.rows[0]._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    properties.append(header)


def _condition_text(condition: ReportCondition) -> str:
    if condition.operator == "is_null":
        return f"{condition.feature.value} IS NULL"
    value = repr(condition.numeric_value) if condition.numeric_value is not None else condition.value_token or "N/A"
    return f"{condition.feature.value} {condition.operator} {value}"


def _rule_text(rule: ReportRule) -> str:
    return " AND ".join(_condition_text(item) for item in rule.conditions) or "N/A"


def _number(value: float | None) -> str:
    return "N/A" if value is None else f"{value:.4f}"


def _integer(value: int | None) -> str:
    return "N/A" if value is None else str(value)


def _distribution(value: MetricDistribution) -> str:
    if value.count == 0:
        return "N/A"
    return f"count={value.count}, min={_number(value.minimum)}, median={_number(value.median)}, max={_number(value.maximum)}"


def _canonicalize_docx_package(content: bytes) -> bytes:
    target_buffer = BytesIO()
    with ZipFile(BytesIO(content), "r") as source:
        names = source.namelist()
        if len(names) != len(set(names)):
            raise ValueError("DOCX package contains duplicate members")
        with ZipFile(
            target_buffer,
            "w",
            compression=ZIP_DEFLATED,
            compresslevel=9,
        ) as target:
            for name in sorted(names):
                source_info = source.getinfo(name)
                info = ZipInfo(filename=name, date_time=_FIXED_ZIP_TIME)
                info.compress_type = ZIP_STORED if source_info.is_dir() else ZIP_DEFLATED
                info.create_system = 3
                info.external_attr = (0o40700 if source_info.is_dir() else 0o100600) << 16
                target.writestr(
                    info,
                    source.read(name),
                    compress_type=info.compress_type,
                    compresslevel=9 if info.compress_type == ZIP_DEFLATED else None,
                )
    return target_buffer.getvalue()


__all__ = ["render_report_docx"]
