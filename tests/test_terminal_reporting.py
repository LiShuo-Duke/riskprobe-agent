import os
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory
from zipfile import ZipFile

import pytest
from docx import Document

from riskprobe.analysis_contracts import (
    AnalysisSummary,
    DecisionSummary,
    StageName,
    StageStatus,
    StageSummary,
)
from riskprobe.agents.contracts import (
    AgentResult,
    AgentState,
    AgentStatus,
    ExecutionPlan,
    PlanStep,
    ReviewDecision,
)
from riskprobe.models import Condition, EvidenceCard, EvidenceGrade, RiskRule, RuleMetrics
from riskprobe.report_models import (
    ReportModel,
    ReportReview,
    build_analysis_report_model,
    build_final_report_model,
)
from riskprobe.reporting import render_report_markdown
from riskprobe.reporting_docx import render_report_docx
from riskprobe.terminal_reports import (
    TerminalReportError,
    TerminalReportStore,
    TerminalReportSubject,
)
from riskprobe.tools import (
    DiagnoseRequest,
    DiscoverRequest,
    InspectRequest,
    RecommendRequest,
)


def _metrics(lift: float) -> RuleMetrics:
    return RuleMetrics(
        support_count=20,
        coverage=0.2,
        base_bad_rate=0.1,
        hit_bad_rate=0.3,
        non_hit_bad_rate=0.05,
        lift=lift,
        precision=0.3,
        recall=0.6,
        p_value=0.01,
        hit_good_rate=0.1,
        ks_signed=0.2,
        ks_stat=0.2,
    )


def _card(index: int, *, grade: EvidenceGrade, lift: float) -> EvidenceCard:
    return EvidenceCard(
        rule=RiskRule(
            rule_id=f"rule-{index:02d}",
            origin="single",
            conditions=(Condition(feature="amount", operator=">", value=float(index)),),
        ),
        train=_metrics(lift - 0.1),
        test=_metrics(lift),
        slices=(),
        lift_ci=(lift - 0.2, lift + 0.2),
        adjusted_p_value=0.02,
        segment_consistency=1.0,
        max_time_decay=0.0,
        grade=grade,
        limitations=(),
    )


def test_report_model_limits_orders_and_redacts_rules() -> None:
    cards = [
        _card(
            index,
            grade="Stable" if index < 2 else "Unstable",
            lift=float(index),
        )
        for index in range(11)
    ]
    cards[0] = cards[0].model_copy(
        update={
            "rule": RiskRule(
                rule_id="rule-00",
                origin="single",
                conditions=(
                    Condition(
                        feature="merchant_type",
                        operator="==",
                        value="private-category",
                    ),
                ),
            )
        }
    )

    first = build_analysis_report_model(
        evidence_cards=cards,
        time_validation_applied=False,
    )
    second = build_analysis_report_model(
        evidence_cards=cards,
        time_validation_applied=False,
    )

    assert len(first.top_rules) == 10
    assert [rule.rule_id for rule in first.top_rules[:2]] == ["rule-01", "rule-00"]
    assert first.top_rules[0].holdout is None
    assert first.top_rules[0].time_decay is None
    assert first.time_validation_applied is False
    assert first.model_dump(mode="json") == second.model_dump(mode="json")
    assert "private-category" not in first.model_dump_json()
    assert first.top_rules[1].conditions[0].value_token.startswith("tok_")

    markdown = render_report_markdown(first)
    assert "## 5. TOP10 规则" in markdown
    assert "规则条件" in markdown
    assert "Test Lift" in markdown
    assert "Holdout Lift" in markdown
    assert "private-category" not in markdown


def test_docx_is_deterministic_openable_and_redacted() -> None:
    card = _card(1, grade="Stable", lift=3.0).model_copy(
        update={
            "rule": RiskRule(
                rule_id="rule-docx",
                origin="single",
                conditions=(
                    Condition(
                        feature="merchant_type",
                        operator="==",
                        value="private-category",
                    ),
                ),
            )
        }
    )
    model = build_analysis_report_model(
        evidence_cards=(card,),
        time_validation_applied=False,
    )

    first = render_report_docx(model)
    second = render_report_docx(model)

    assert first == second
    with ZipFile(BytesIO(first), "r") as package:
        document_xml = package.read("word/document.xml")
        core_xml = package.read("docProps/core.xml")
        document_text = document_xml.decode("utf-8")
        assert b"rule-docx" in document_xml
        assert b"private-category" not in document_xml
        assert b"/Users/" not in document_xml
        assert "时间切片稳定性：未评估" in document_text
        assert "Time Decay" not in document_text
        assert b"test@example.com" not in core_xml
    assert Document(BytesIO(first)).paragraphs


def test_no_action_renders_message_and_review_in_both_formats() -> None:
    model = ReportModel(
        scope="terminal",
        terminal_status="no_action",
        review=ReportReview(
            approved=True,
            retry_allowed=False,
            no_action_required=True,
        ),
    )

    markdown = render_report_markdown(model)
    assert "无需 Host proposal" in markdown
    assert "### Review" in markdown
    assert "No action required: True" in markdown

    document = Document(BytesIO(render_report_docx(model)))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "无需 Host proposal" in text
    assert "Review" in text
    assert "No action required: True" in text


def test_terminal_report_store_reuses_repairs_and_rejects_tampering() -> None:
    model = ReportModel(
        scope="terminal",
        run_id="run-check",
        terminal_status="failed",
        error_code="agent_orchestration_failed",
    )
    subject = TerminalReportSubject(
        idempotency_key="terminal-report-check",
        run_id="run-check",
        context_id=None,
        findings=(),
        proposal_action_codes=(),
        diagnosis_evidence_ids=(),
        agent_result=None,
        analysis_summary=None,
        decision_summary=None,
        terminal_status="failed",
        error_code="agent_orchestration_failed",
    )

    with TemporaryDirectory() as directory:
        root = Path(directory)
        store = TerminalReportStore(root)
        try:
            store.ensure_published(
                subject=subject,
                model=model.model_copy(update={"run_id": "other-run"}),
            )
        except TerminalReportError as error:
            assert error.code == "terminal_report_integrity_failed"
        else:
            raise AssertionError("mismatched terminal report binding must fail closed")

        first = store.ensure_published(subject=subject, model=model)
        second = store.ensure_published(subject=subject, model=model)
        assert first == second

        report_dir = root / ".riskprobe-terminal-reports" / first.report_id
        markdown_path = report_dir / "final_risk_report.md"
        docx_path = report_dir / "final_risk_report.docx"
        manifest_path = report_dir / "final_report_manifest.json"
        assert {item.name for item in report_dir.iterdir() if not item.name.endswith(".lock")} == {
            markdown_path.name,
            docx_path.name,
            manifest_path.name,
        }
        assert os.stat(markdown_path).st_mode & 0o777 == 0o600
        assert os.stat(docx_path).st_mode & 0o777 == 0o600

        original_docx_hash = first.docx.sha256
        manifest_path.unlink()
        repaired = store.ensure_published(subject=subject, model=model)
        assert repaired.docx.sha256 == original_docx_hash

        docx_path.chmod(0o600)
        docx_path.write_bytes(b"tampered")
        try:
            store.ensure_published(subject=subject, model=model)
        except TerminalReportError as error:
            assert error.code == "terminal_report_integrity_failed"
        else:
            raise AssertionError("tampered terminal report must fail closed")


def _summary_with_stages(
    updates: dict[StageName, StageSummary],
) -> AnalysisSummary:
    return AnalysisSummary(
        stages=tuple(
            updates.get(
                name,
                StageSummary(
                    name=name,
                    status=StageStatus.NOT_RUN,
                    enabled=True,
                    output_available=False,
                ),
            )
            for name in StageName
        )
    )


def _no_action_agent_result(
    analysis_summary: AnalysisSummary | None,
) -> AgentResult:
    return AgentResult(
        session_id="terminal-no-action-run",
        status=AgentStatus.SUCCEEDED,
        plan=ExecutionPlan(
            objective="comprehensive",
            dataset_id="synthetic_demo",
            steps=(
                PlanStep(
                    step_id="inspect",
                    tool_name="inspect",
                    request=InspectRequest(dataset_id="synthetic_demo"),
                ),
                PlanStep(
                    step_id="diagnose",
                    tool_name="diagnose",
                    request=DiagnoseRequest(dataset_id="synthetic_demo"),
                ),
                PlanStep(
                    step_id="discover",
                    tool_name="discover",
                    request=DiscoverRequest(dataset_id="synthetic_demo"),
                ),
                PlanStep(
                    step_id="recommend",
                    tool_name="recommend",
                    request=RecommendRequest(dataset_id="synthetic_demo"),
                ),
                PlanStep(step_id="review", tool_name="review"),
            ),
            component_versions={"planner": "planner-v1"},
        ),
        review=ReviewDecision(approved=True, no_action_required=True),
        tool_sequence=("inspect", "diagnose", "discover", "recommend", "review"),
        retry_count=0,
        state_history=(
            AgentState.PLANNING,
            AgentState.EXECUTING,
            AgentState.COLLECTING_EVIDENCE,
            AgentState.REVIEWING,
            AgentState.COMPLETED,
        ),
        leaf_node_id="d" * 64,
        redacted_summary="no actionable diagnosis",
        analysis_summary=analysis_summary,
    )


def test_no_action_report_consumes_terminal_outcome_summary_exactly() -> None:
    partial = _summary_with_stages(
        {
            StageName.INSPECT: StageSummary(
                name=StageName.INSPECT,
                status=StageStatus.SUCCEEDED,
                enabled=True,
                output_available=True,
            ),
            StageName.DIAGNOSE_QUALITY: StageSummary(
                name=StageName.DIAGNOSE_QUALITY,
                status=StageStatus.UNAVAILABLE,
                enabled=True,
                output_available=False,
                reason_code="diagnostic_profile_unavailable",
            ),
            StageName.DIAGNOSE_RULE_EVIDENCE: StageSummary(
                name=StageName.DIAGNOSE_RULE_EVIDENCE,
                status=StageStatus.NOT_RUN,
                enabled=False,
                output_available=False,
                reason_code="rule_evidence_not_run",
            ),
            StageName.DISCOVER_RESTORE: StageSummary(
                name=StageName.DISCOVER_RESTORE,
                status=StageStatus.SUCCEEDED,
                enabled=True,
                output_available=True,
            ),
        }
    )
    terminal = _summary_with_stages(
        {
            **{stage.name: stage for stage in partial.stages},
            StageName.DECISION_CONTEXT: StageSummary(
                name=StageName.DECISION_CONTEXT,
                status=StageStatus.SKIPPED,
                enabled=True,
                output_available=False,
                reason_code="no_action_required",
            ),
            StageName.RECOMMEND: StageSummary(
                name=StageName.RECOMMEND,
                status=StageStatus.SKIPPED,
                enabled=True,
                output_available=False,
                reason_code="no_action_required",
            ),
            StageName.REVIEW: StageSummary(
                name=StageName.REVIEW,
                status=StageStatus.SUCCEEDED,
                enabled=True,
                output_available=True,
            ),
            StageName.TERMINAL: StageSummary(
                name=StageName.TERMINAL,
                status=StageStatus.SUCCEEDED,
                enabled=True,
                output_available=True,
            ),
        }
    )
    result = _no_action_agent_result(partial)
    decision = DecisionSummary(
        selected_action_codes=(),
        recommendations=(),
        recommendation_status=StageStatus.SKIPPED,
        review_approved=True,
        review_reason_codes=(),
        no_action_required=True,
        retry_count=0,
        tool_sequence=result.tool_sequence,
        evidence_complete=True,
        decision_status="no_action",
        final_status="succeeded",
    )
    subject = TerminalReportSubject(
        idempotency_key="terminal-no-action",
        run_id=result.session_id,
        context_id=None,
        findings=(),
        proposal_action_codes=(),
        diagnosis_evidence_ids=(),
        agent_result=result,
        analysis_summary=terminal,
        decision_summary=decision,
        terminal_status="no_action",
    )

    report = build_final_report_model(
        subject=subject,
        evidence_cards=(),
        artifact_analysis_summary=_summary_with_stages({}),
        time_validation_applied=False,
    )

    assert report.analysis_summary == terminal
    assert report.decision_summary == decision
    stages = {stage.name: stage for stage in report.analysis_summary.stages}
    terminal_stages = {stage.name: stage for stage in terminal.stages}
    for name in (
        StageName.DIAGNOSE_QUALITY,
        StageName.DIAGNOSE_FEATURE_DRIFT,
        StageName.DIAGNOSE_POPULATION_SHIFT,
        StageName.DIAGNOSE_TARGET_SHIFT,
        StageName.DIAGNOSE_SEGMENT_RISK,
        StageName.DIAGNOSE_TIME_STABILITY,
        StageName.DIAGNOSE_RULE_EVIDENCE,
    ):
        assert stages[name] == terminal_stages[name]
    assert stages[StageName.DIAGNOSE_QUALITY].status is StageStatus.UNAVAILABLE
    assert stages[StageName.DIAGNOSE_RULE_EVIDENCE].status is StageStatus.NOT_RUN


def test_legacy_no_action_report_does_not_fabricate_missing_summaries() -> None:
    result = _no_action_agent_result(None)
    subject = TerminalReportSubject(
        idempotency_key="legacy-terminal-no-action",
        run_id=result.session_id,
        context_id=None,
        findings=(),
        proposal_action_codes=(),
        diagnosis_evidence_ids=(),
        agent_result=result,
        analysis_summary=None,
        decision_summary=None,
        terminal_status="no_action",
    )

    with pytest.raises(ValueError, match="terminal analysis summary"):
        build_final_report_model(
            subject=subject,
            evidence_cards=(),
            artifact_analysis_summary=_summary_with_stages({}),
            time_validation_applied=False,
        )
